"""Workflow: extract searchable text from downloaded Slack attachments."""

from __future__ import annotations

import hashlib
import io
import os
import zipfile
from dataclasses import dataclass, field
from pathlib import PurePath
from typing import Any

from api.runtime_control import canonical_json
from api.workflow_engine import WorkflowContext

from workflows.slack.shared import env_flag_enabled, positive_int

WORKFLOW_NAME = "slack_attachment_text_extraction"

EXTRACTOR_VERSION = "1"
DEFAULT_INTERVAL_SECONDS = 5 * 60
DEFAULT_BATCH_SIZE = 10
DEFAULT_MAX_EXPANDED_BYTES = 32 * 1024 * 1024
DEFAULT_MAX_PAGES = 250
DEFAULT_MAX_CHARACTERS = 2_000_000

PDF_MIME_TYPE = "application/pdf"
DOCX_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
PPTX_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument.presentationml.presentation"
)
TEXT_MIME_TYPES = {"text/plain", "text/markdown", "text/x-markdown"}

SCHEDULE = {
    "schedule_id": WORKFLOW_NAME,
    "interval_seconds": positive_int(
        os.getenv("SLACK_ATTACHMENT_EXTRACTION_INTERVAL_SECONDS"),
        DEFAULT_INTERVAL_SECONDS,
    ),
    "enabled": (
        env_flag_enabled("SLACK_ETL_ENABLED", default=False)
        and env_flag_enabled("SLACK_ATTACHMENT_EXTRACTION_ENABLED", default=True)
    ),
    "no_delivery": True,
}


@dataclass
class Input:
    """Runtime options for one attachment extraction batch."""

    batch_size: int | None = None
    max_expanded_bytes: int | None = None
    max_pages: int | None = None
    max_characters: int | None = None
    metadata: dict[str, Any] = field(default_factory=dict)


class UnsupportedAttachment(ValueError):
    """The attachment cannot produce searchable text in this extractor."""

    def __init__(self, reason: str, metadata: dict[str, Any] | None = None) -> None:
        super().__init__(reason)
        self.reason = reason
        self.metadata = metadata or {}


def _configured_positive_int(
    explicit: int | None,
    env_name: str,
    default: int,
) -> int:
    return positive_int(
        explicit if explicit is not None else os.getenv(env_name), default
    )


def _suffix(filename: str) -> str:
    return PurePath(filename.lower()).suffix


def _detected_format(data: bytes, filename: str, declared_mime_type: str) -> str:
    mime_type = declared_mime_type.lower().split(";", 1)[0].strip()
    suffix = _suffix(filename)
    if data.startswith(b"%PDF-"):
        return "pdf"
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                names = set(archive.namelist())
        except zipfile.BadZipFile as error:
            raise UnsupportedAttachment("invalid_zip_container") from error
        if "word/document.xml" in names:
            return "docx"
        if "ppt/presentation.xml" in names:
            return "pptx"
    if mime_type == PDF_MIME_TYPE or suffix == ".pdf":
        return "pdf"
    if mime_type == DOCX_MIME_TYPE or suffix == ".docx":
        return "docx"
    if mime_type == PPTX_MIME_TYPE or suffix == ".pptx":
        return "pptx"
    if mime_type in TEXT_MIME_TYPES or suffix in {".txt", ".md", ".markdown"}:
        return "text"
    raise UnsupportedAttachment("unsupported_file_type")


def _check_expanded_size(data: bytes, max_expanded_bytes: int) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            expanded_bytes = sum(entry.file_size for entry in archive.infolist())
    except zipfile.BadZipFile as error:
        raise UnsupportedAttachment("invalid_zip_container") from error
    if expanded_bytes > max_expanded_bytes:
        raise UnsupportedAttachment(
            "expanded_file_too_large",
            {
                "expanded_bytes": expanded_bytes,
                "max_expanded_bytes": max_expanded_bytes,
            },
        )


def _truncate(text: str, max_characters: int) -> tuple[str, bool]:
    normalized = text.replace("\x00", "").strip()
    if len(normalized) <= max_characters:
        return normalized, False
    return normalized[:max_characters].rstrip(), True


def _extract_pdf(data: bytes, max_pages: int) -> tuple[str, dict[str, Any]]:
    try:
        import pymupdf
    except ImportError as error:  # pragma: no cover - deployment packaging guard
        raise RuntimeError("PyMuPDF is required for PDF extraction") from error

    try:
        document = pymupdf.open(stream=data, filetype="pdf")
    except Exception as error:
        raise UnsupportedAttachment("invalid_or_encrypted_pdf") from error

    try:
        if document.needs_pass:
            raise UnsupportedAttachment(
                "password_protected_pdf", {"password_protected": True}
            )
        page_count = document.page_count
        processed_pages = min(page_count, max_pages)
        pages_without_text: list[int] = []
        sections: list[str] = []
        for page_index in range(processed_pages):
            text = document.load_page(page_index).get_text("text").strip()
            if not text:
                pages_without_text.append(page_index + 1)
                continue
            sections.append(f"## Page {page_index + 1}\n\n{text}")
    finally:
        document.close()

    metadata = {
        "detected_mime_type": PDF_MIME_TYPE,
        "page_count": page_count,
        "processed_pages": processed_pages,
        "pages_without_text": pages_without_text,
        "page_limit_reached": page_count > max_pages,
    }
    if not sections:
        raise UnsupportedAttachment("no_extractable_text", metadata)
    return "\n\n".join(sections), metadata


def _extract_docx(
    data: bytes,
    max_expanded_bytes: int,
) -> tuple[str, dict[str, Any]]:
    _check_expanded_size(data, max_expanded_bytes)
    try:
        from docx import Document
    except ImportError as error:  # pragma: no cover - deployment packaging guard
        raise RuntimeError("python-docx is required for DOCX extraction") from error

    try:
        document = Document(io.BytesIO(data))
    except Exception as error:
        raise UnsupportedAttachment("invalid_or_encrypted_docx") from error

    sections = [paragraph.text.strip() for paragraph in document.paragraphs]
    for table in document.tables:
        rows = [
            "\t".join(cell.text.strip() for cell in row.cells).strip()
            for row in table.rows
        ]
        sections.extend(row for row in rows if row)
    text = "\n\n".join(section for section in sections if section)
    if not text:
        raise UnsupportedAttachment("no_extractable_text")
    return text, {
        "detected_mime_type": DOCX_MIME_TYPE,
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
    }


def _extract_pptx(
    data: bytes,
    max_expanded_bytes: int,
) -> tuple[str, dict[str, Any]]:
    _check_expanded_size(data, max_expanded_bytes)
    try:
        from pptx import Presentation
    except ImportError as error:  # pragma: no cover - deployment packaging guard
        raise RuntimeError("python-pptx is required for PPTX extraction") from error

    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as error:
        raise UnsupportedAttachment("invalid_or_encrypted_pptx") from error

    slides: list[str] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = str(shape.text or "").strip()
                if text:
                    parts.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    text = "\t".join(cell.text.strip() for cell in row.cells).strip()
                    if text:
                        parts.append(text)
        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame
            notes = str(notes_frame.text or "").strip() if notes_frame else ""
            if notes:
                parts.append(f"Speaker notes:\n{notes}")
        if parts:
            slides.append(f"## Slide {slide_number}\n\n" + "\n\n".join(parts))
    if not slides:
        raise UnsupportedAttachment("no_extractable_text")
    return "\n\n".join(slides), {
        "detected_mime_type": PPTX_MIME_TYPE,
        "slide_count": len(presentation.slides),
    }


def extract_text(
    data: bytes,
    *,
    filename: str,
    declared_mime_type: str,
    max_expanded_bytes: int,
    max_pages: int,
    max_characters: int,
) -> tuple[str, dict[str, Any]]:
    """Extract bounded searchable text and metadata from one attachment."""
    detected_format = _detected_format(data, filename, declared_mime_type)
    if detected_format == "pdf":
        text, metadata = _extract_pdf(data, max_pages)
    elif detected_format == "docx":
        text, metadata = _extract_docx(data, max_expanded_bytes)
    elif detected_format == "pptx":
        text, metadata = _extract_pptx(data, max_expanded_bytes)
    else:
        text = data.decode("utf-8", errors="replace")
        metadata = {"detected_mime_type": "text/plain"}
        if not text.strip():
            raise UnsupportedAttachment("no_extractable_text", metadata)

    text, truncated = _truncate(text, max_characters)
    metadata = {
        **metadata,
        "character_count": len(text),
        "truncated": truncated or bool(metadata.get("page_limit_reached")),
    }
    return text, metadata


async def _load_candidates(pool, *, batch_size: int) -> list[Any]:
    return list(
        await pool.fetch(
            "SELECT a.channel_id, a.message_ts, a.slack_file_id, a.name, a.mimetype, "
            "a.content_sha256, a.content_bytes, a.updated_at "
            "FROM slack_sync_message_attachments a "
            "LEFT JOIN slack_attachment_text_extractions e "
            "ON e.channel_id = a.channel_id "
            "AND e.message_ts = a.message_ts "
            "AND e.slack_file_id = a.slack_file_id "
            "WHERE a.download_status = 'downloaded' "
            "AND a.content_bytes IS NOT NULL "
            "AND ("
            "  lower(a.mimetype) IN ($1, $2, $3, 'text/plain', 'text/markdown', 'text/x-markdown') "
            "  OR lower(a.filetype) IN ('pdf', 'docx', 'pptx', 'txt', 'text', 'md', 'markdown') "
            "  OR lower(a.name) ~ '\\.(pdf|docx|pptx|txt|md|markdown)$'"
            ") "
            "AND (e.extraction_id IS NULL "
            "  OR e.extractor_version IS DISTINCT FROM $4 "
            "  OR (a.content_sha256 IS NOT NULL "
            "    AND e.source_content_sha256 IS DISTINCT FROM a.content_sha256)) "
            "ORDER BY a.updated_at, a.channel_id, a.message_ts, a.slack_file_id "
            "LIMIT $5",
            PDF_MIME_TYPE,
            DOCX_MIME_TYPE,
            PPTX_MIME_TYPE,
            EXTRACTOR_VERSION,
            batch_size,
        )
    )


async def _store_result(
    pool,
    *,
    row: Any,
    source_content_sha256: str | None,
    status: str,
    text_content: str = "",
    metadata: dict[str, Any] | None = None,
    last_error: str = "",
) -> None:
    # Keep the extraction and the source-row freshness signal atomic. The
    # company-context attachment projection checkpoints the source updated_at.
    await pool.execute(
        "WITH extraction AS ("
        "  INSERT INTO slack_attachment_text_extractions ("
        "    channel_id, message_ts, slack_file_id, source_content_sha256, "
        "    extractor_version, status, text_content, metadata, last_error"
        "  ) VALUES ($1, $2, $3, $4, $5, $6, $7, $8::jsonb, $9) "
        "  ON CONFLICT (channel_id, message_ts, slack_file_id) DO UPDATE SET "
        "    source_content_sha256 = EXCLUDED.source_content_sha256, "
        "    extractor_version = EXCLUDED.extractor_version, "
        "    status = EXCLUDED.status, "
        "    text_content = EXCLUDED.text_content, "
        "    metadata = EXCLUDED.metadata, "
        "    last_error = EXCLUDED.last_error, "
        "    updated_at = NOW() "
        "  RETURNING 1"
        ") "
        "UPDATE slack_sync_message_attachments SET updated_at = NOW() "
        "WHERE channel_id = $1 AND message_ts = $2 AND slack_file_id = $3 "
        "AND EXISTS (SELECT 1 FROM extraction)",
        row["channel_id"],
        row["message_ts"],
        row["slack_file_id"],
        source_content_sha256,
        EXTRACTOR_VERSION,
        status,
        text_content,
        canonical_json(metadata or {}),
        last_error[:1_000],
    )


async def _extract_and_store(
    pool,
    row: Any,
    *,
    max_expanded_bytes: int,
    max_pages: int,
    max_characters: int,
) -> dict[str, Any]:
    data = bytes(row["content_bytes"])
    source_hash = hashlib.sha256(data).hexdigest()
    try:
        text, metadata = extract_text(
            data,
            filename=str(row["name"] or ""),
            declared_mime_type=str(row["mimetype"] or ""),
            max_expanded_bytes=max_expanded_bytes,
            max_pages=max_pages,
            max_characters=max_characters,
        )
    except UnsupportedAttachment as error:
        await _store_result(
            pool,
            row=row,
            source_content_sha256=source_hash,
            status="unsupported",
            metadata=error.metadata,
            last_error=error.reason,
        )
        return {"status": "unsupported", "reason": error.reason}
    # Third-party parsers expose unrelated exception hierarchies. Persist their
    # failures so a malformed attachment does not block the remaining batch.
    except Exception as error:  # noqa: BLE001
        await _store_result(
            pool,
            row=row,
            source_content_sha256=source_hash,
            status="failed",
            last_error=f"{type(error).__name__}: {error}",
        )
        return {"status": "failed", "error_type": type(error).__name__}

    await _store_result(
        pool,
        row=row,
        source_content_sha256=source_hash,
        status="succeeded",
        text_content=text,
        metadata=metadata,
    )
    return {"status": "succeeded", "characters": len(text)}


def _step_name(row: Any) -> str:
    identity = ":".join(
        [
            str(row["channel_id"]),
            str(row["message_ts"]),
            str(row["slack_file_id"]),
            str(row["content_sha256"] or row["updated_at"] or "unknown"),
            EXTRACTOR_VERSION,
        ]
    )
    return f"extract:{hashlib.sha256(identity.encode('utf-8')).hexdigest()}"


async def handler(inp: Input, ctx: WorkflowContext) -> dict[str, Any]:
    """Extract one bounded batch and requeue while more candidates remain."""
    if not env_flag_enabled("SLACK_ATTACHMENT_EXTRACTION_ENABLED", default=True):
        return {"status": "skipped", "reason": "attachment_extraction_disabled"}

    batch_size = _configured_positive_int(
        inp.batch_size,
        "SLACK_ATTACHMENT_EXTRACTION_BATCH_SIZE",
        DEFAULT_BATCH_SIZE,
    )
    max_expanded_bytes = _configured_positive_int(
        inp.max_expanded_bytes,
        "SLACK_ATTACHMENT_EXTRACTION_MAX_EXPANDED_BYTES",
        DEFAULT_MAX_EXPANDED_BYTES,
    )
    max_pages = _configured_positive_int(
        inp.max_pages,
        "SLACK_ATTACHMENT_EXTRACTION_MAX_PAGES",
        DEFAULT_MAX_PAGES,
    )
    max_characters = _configured_positive_int(
        inp.max_characters,
        "SLACK_ATTACHMENT_EXTRACTION_MAX_CHARACTERS",
        DEFAULT_MAX_CHARACTERS,
    )
    rows = await _load_candidates(ctx._pool, batch_size=batch_size)

    counts = {"succeeded": 0, "unsupported": 0, "failed": 0}
    for row in rows:
        result = await ctx.step(
            _step_name(row),
            lambda row=row: _extract_and_store(
                ctx._pool,
                row,
                max_expanded_bytes=max_expanded_bytes,
                max_pages=max_pages,
                max_characters=max_characters,
            ),
            step_kind="slack_attachment_text_extraction",
        )
        counts[str(result["status"])] += 1

    next_run = None
    if len(rows) == batch_size:
        next_run = await ctx.start_workflow(
            WORKFLOW_NAME,
            {
                "batch_size": batch_size,
                "max_expanded_bytes": max_expanded_bytes,
                "max_pages": max_pages,
                "max_characters": max_characters,
                "metadata": {
                    **inp.metadata,
                    "source": "slack_attachment_text_extraction_requeue",
                },
            },
            idempotency_key=f"{WORKFLOW_NAME}:{ctx.run_id}:next",
        )

    result: dict[str, Any] = {
        "status": "completed",
        **counts,
        "extractor_version": EXTRACTOR_VERSION,
        "requeued": next_run is not None,
    }
    if next_run is not None:
        result["next_run"] = next_run
    ctx.log("slack_attachment_text_extraction_completed", **result)
    return result
